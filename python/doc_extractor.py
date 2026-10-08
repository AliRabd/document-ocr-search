#!/usr/bin/env python3
"""DOC / DOCX / XLSX / CSV / TXT / text-layer PDF -> JSON Lines (chunked)."""
import csv
import os
import subprocess
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common  # noqa: E402
from common import MAX_TOTAL_CHARS, make_records  # noqa: E402

CMD_TIMEOUT = int(os.environ.get("CMD_TIMEOUT", "300"))


def read_text_file(path):
    last = None
    for enc in common.candidate_encodings(path):
        try:
            with open(path, "r", encoding=enc, newline="") as handle:
                return handle.read(MAX_TOTAL_CHARS + 1)
        except (UnicodeError, LookupError) as exc:
            last = exc
    raise ValueError("cannot decode file: %s" % last)


def extract_csv(path):
    csv.field_size_limit(2 ** 27)
    for enc in common.candidate_encodings(path):
        try:
            with open(path, "r", encoding=enc, newline="") as handle:
                sample = handle.read(4096)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                lines, total = [], 0
                for row in csv.reader(handle, dialect):
                    line = ", ".join(cell.strip() for cell in row)
                    lines.append(line)
                    total += len(line) + 1
                    if total > MAX_TOTAL_CHARS:
                        break
                return "\n".join(lines)
        except (UnicodeError, LookupError):
            continue
    raise ValueError("cannot decode CSV with utf-8 / utf-16 / cp1256")


def extract_xlsx(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts, total, cut = [], 0, False
    try:
        for sheet in workbook.worksheets:           # ALL sheets, not just the first
            parts.append("## " + sheet.title)
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row
                         if c is not None and str(c).strip()]
                if cells:
                    line = ", ".join(cells)
                    parts.append(line)
                    total += len(line) + 1
                if total > MAX_TOTAL_CHARS:
                    cut = True
                    break
            if cut:
                break
        return "\n".join(parts), len(workbook.worksheets), cut
    finally:
        workbook.close()


def run_tool(cmd):
    out = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                         timeout=CMD_TIMEOUT, check=True).stdout
    return out.decode("utf-8", "replace")


def extract(path):
    """Return (text, doc_type, engine, extra)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        text = run_tool(["pdftotext", "-enc", "UTF-8", path, "-"])
        return text, "pdf_text", "pdftotext", {"page_count": text.count("\f")}
    if ext == ".doc":
        try:
            text = run_tool(["antiword", "-m", "UTF-8.txt", path])
        except subprocess.CalledProcessError:
            text = run_tool(["antiword", path])
        return text, "word", "antiword", None
    if ext == ".docx":
        import docx2txt
        return docx2txt.process(path), "word", "docx2txt", None
    if ext in (".xlsx", ".xlsm"):
        text, sheets, cut = extract_xlsx(path)
        return text, "excel", "openpyxl", {"page_count": sheets, "_truncated": cut}
    if ext == ".csv":
        return extract_csv(path), "csv", "csv", None
    if ext in (".txt", ".md"):
        return read_text_file(path), "text", "plain", None
    return "", "unsupported", "none", None


def main(argv):
    common.setup_stdio()
    if not argv:
        common.log("Usage: doc_extractor.py <file> [<file> ...]")
        return 1
    for path in argv:
        start = time.time()
        try:
            text, doc_type, engine, extra = extract(path)
            extra = dict(extra or {})
            truncated = bool(extra.pop("_truncated", False))
            records = make_records(path, text, doc_type, engine, extra,
                                   truncated=truncated,
                                   elapsed=time.time() - start)
        except Exception as exc:
            msg = "Failed to process %s: %s" % (os.path.basename(path), exc)
            common.log(msg)
            records = [common.error_record(path, "document", msg)]
        common.print_records(records)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
